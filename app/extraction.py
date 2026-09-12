import io
import re
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.config import settings


def extract(data: bytes, filename: str) -> list[dict]:
    """Parse bounded input without writing originals to disk or following external relationships."""
    if not data:
        raise ValueError("This file is empty.")
    if len(data) > settings.max_file_bytes:
        raise ValueError("Files must be 10 MB or smaller.")
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ValueError("This does not appear to be a PDF file.")
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")
            if len(reader.pages) > settings.max_pages:
                raise ValueError(f"PDFs must have at most {settings.max_pages} pages.")
            # Bound decompressed page content as well as the compressed upload.
            sections = []
            total = 0
            for number, page in enumerate(reader.pages, 1):
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > 2_000_000:
                    raise ValueError("A PDF page is too complex for this demo.")
                text = page.extract_text() or ""
                total += len(text)
                if total > settings.max_text_chars:
                    raise ValueError("The document contains too much text for this demo.")
                sections.append({"location": f"Page {number}", "text": text})
        elif extension == ".docx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 20_000_000:
                    raise ValueError("The expanded DOCX is too large for this demo.")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("This does not appear to be a DOCX file.")
            document = Document(io.BytesIO(data))
            blocks = [p.text for p in document.paragraphs]
            blocks += [" | ".join(cell.text for cell in row.cells)
                       for table in document.tables for row in table.rows]
            # DOCX does not have stable page numbers without a layout engine.
            sections = [{"location": f"Paragraph {i}", "text": text}
                        for i, text in enumerate(blocks, 1) if text.strip()]
        elif extension == ".txt":
            text = data.decode("utf-8-sig")
            sections = [{"location": f"Section {i}", "text": text}
                        for i, text in enumerate(re.split(r"\n\s*\n", text), 1) if text.strip()]
        else:
            raise ValueError("Supported file types are PDF, DOCX, and UTF-8 TXT.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("The file could not be read. Check its format and try another file.") from exc
    if sum(len(s["text"]) for s in sections) > settings.max_text_chars:
        raise ValueError("The document contains too much text (150,000 characters maximum).")
    sections = [{**s, "text": s["text"].replace("\x00", "").strip()}
                for s in sections if s["text"].strip()]
    if not sections:
        raise ValueError("No readable text was found. Scanned PDFs need OCR, which this demo excludes.")
    return sections


def chunks(sections: list[dict], size: int = 1000, overlap: int = 150) -> list[dict]:
    result = []
    for section in sections:
        text = section["text"]
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                boundary = text.rfind(" ", start + size // 2, end)
                if boundary > start:
                    end = boundary
            result.append({"location": section["location"], "text": text[start:end].strip()})
            if len(result) > settings.max_chunks:
                raise ValueError("The document has too many passages for this demo.")
            if end == len(text):
                break
            start = max(start + 1, end - overlap)
    return result
