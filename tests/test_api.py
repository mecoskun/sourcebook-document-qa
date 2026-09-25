import io
import time

from docx import Document
from fastapi.testclient import TestClient

from app import main


def test_auth_expiry_and_input_limits():
    # Exercise authentication and validation against a real MCP process without model downloads.
    with TestClient(main.app, raise_server_exceptions=False) as client:
        assert client.get("/api/documents").status_code == 401
        response = client.post("/api/sessions")
        token = response.json()["token"]
        headers = {"Authorization": "Bearer " + token}
        assert client.get("/api/documents", headers=headers).json() == []
        assert client.post("/api/questions", headers=headers,
                           json={"question": "hello"}).status_code == 400
        assert client.post("/api/questions", headers=headers,
                           json={"question": "x" * 1001}).status_code == 422
        assert client.post("/api/questions", headers=headers,
                           content=b"x" * 9000).status_code == 413
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Ambiguous merged policy"
        buffer = io.BytesIO()
        document.save(buffer)
        rejected = client.post("/api/documents", headers=headers,
                               files={"file": ("merged.docx", buffer.getvalue())})
        assert rejected.status_code == 400
        assert "merged or nested" in rejected.json()["detail"]
        assert client.get("/api/documents", headers=headers).json() == []
        for workspace in main.sessions.values():
            workspace.expires = time.monotonic() - 1
        assert client.get("/api/documents", headers=headers).status_code == 401
    main.sessions.clear()
