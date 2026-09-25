import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.config import settings
from app.extraction import chunks, extract


def test_text_sections_and_chunk_provenance():
    sections = extract(b"First policy.\n\nSecond policy.", "policy.txt")
    assert [s["location"] for s in sections] == ["Section 1", "Section 2"]
    text = "The policy applies to all employees. " * 90
    pieces = chunks([{"location": "Page 7", "text": text}])
    assert len(pieces) > 1
    assert all(p["location"] == "Page 7" and len(p["text"]) <= 1000 for p in pieces)
    assert pieces[-1]["text"].endswith("employees.")


@pytest.mark.parametrize("data,name", [(b"", "empty.txt"), (b"hello", "file.exe"),
    (b"not a PDF", "file.pdf"), (b"\xff\xfe", "file.txt"), (b"bad zip", "file.docx")])
def test_invalid_uploads(data, name):
    with pytest.raises(ValueError):
        extract(data, name)


def test_limits(monkeypatch):
    monkeypatch.setattr(settings, "max_file_bytes", 4)
    with pytest.raises(ValueError, match="10 MB"):
        extract(b"hello", "file.txt")


def test_scanned_and_encrypted_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="No readable text"):
        extract(output.getvalue(), "scan.pdf")
    writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="Password-protected"):
        extract(output.getvalue(), "locked.pdf")


def test_docx_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Annual leave is 20 days.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Carryover"
    table.cell(0, 1).text = "5 days"
    output = io.BytesIO()
    document.save(output)
    sections = extract(output.getvalue(), "policy.docx")
    assert "Annual leave" in sections[0]["text"]
    assert "Carryover | 5 days" in sections[1]["text"]


def test_docx_table_header_context_and_body_order():
    document = Document()
    document.add_paragraph("Service plans")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Plan", "Response time"), ("Basic", "48 hours"), ("Priority", "4 hours")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("Contact the service desk.")
    output = io.BytesIO()
    document.save(output)
    sections = extract(output.getvalue(), "plans.docx")
    assert [s["location"] for s in sections] == ["Paragraph 1", "Table 1, row 2", "Table 1, row 3", "Paragraph 2"]
    assert sections[2]["text"] == "Plan | Response time\nPriority | 4 hours"


@pytest.mark.parametrize("kind", ["horizontal", "vertical", "nested"])
def test_rejects_complex_docx_tables_instead_of_silently_losing_context(kind):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    if kind == "horizontal":
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Combined heading"
    elif kind == "vertical":
        table.cell(0, 0).merge(table.cell(1, 0)).text = "Combined category"
    else:
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Critical nested policy"
    output = io.BytesIO()
    document.save(output)
    with pytest.raises(ValueError, match="merged or nested"):
        extract(output.getvalue(), "complex.docx")
