"""Pre-hosting acceptance suite. Uses one pinned public PDF and controlled edge cases."""

import argparse
import hashlib
import io
import json
import re
import time
from pathlib import Path

import httpx
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PUBLIC_URL = "https://nvlpubs.nist.gov/nistpubs/SpecialPublications/NIST.SP.1300.pdf"
PUBLIC_SHA = "1e77c1bd63d1e15c48c870b7c17b7c9a2edf6bd5103b0940e72828515c9f1881"


def groups(pdf):
    yield (
        "public_pdf",
        [("NIST.SP.1300.pdf", pdf)],
        [
            (
                "List all six CSF Functions.",
                ["govern", "identify", "protect", "detect", "respond", "recover"],
                ["Page 2"],
                False,
            ),
            (
                "What does the guide recommend doing with data backups?",
                ["back", "test"],
                ["Page 5"],
                False,
            ),
            (
                "Give two common indicators of a cybersecurity incident.",
                ["login|access|sluggish|antivirus|traffic|email"],
                ["Page 6"],
                False,
            ),
            (
                "What should be checked before backed-up data is used for restoration?",
                ["integrity"],
                ["Page 8"],
                False,
            ),
            (
                "How does a Current Profile differ from a Target Profile?",
                ["current", "target"],
                ["Page 9"],
                False,
            ),
            ("How much does the cybersecurity insurance policy cost per month?", [], [], True),
            (
                "What should we do with backups regularly, and what should we check before restoring them?",
                ["test", "integrity"],
                ["Page 5", "Page 8"],
                False,
            ),
        ],
    )
    yield (
        "conflicting_sources",
        [
            (
                "alpha.txt",
                b"Company vacation policy: Full-time employees receive 20 paid vacation days per year.",
            ),
            (
                "beta.txt",
                b"Company vacation policy: Full-time employees receive 25 paid vacation days per year.",
            ),
        ],
        [
            (
                "How many paid vacation days do full-time employees receive?",
                ["20", "25", "conflict|differ|inconsisten|disagree"],
                [],
                True,
            ),
            (
                "According to alpha.txt, how many vacation days do full-time employees receive?",
                ["20"],
                [],
                False,
            ),
        ],
    )
    doc = Document()
    doc.add_heading("Operations manual - fictional test fixture", 0)
    for i in range(70):
        doc.add_paragraph(
            f"Station {i + 1}: record the inspection date, inspector name, and condition in the equipment log. Archive completed logs with the operations coordinator."
        )
    doc.add_heading("Laboratory access", 1)
    doc.add_paragraph(
        "Laboratory visitors require an escort. Visitor badges expire after 72 hours. The laboratory opens at 8 AM on weekdays."
    )
    buffer = io.BytesIO()
    doc.save(buffer)
    yield (
        "long_docx",
        [("operations.docx", buffer.getvalue())],
        [
            (
                "How long are laboratory visitor badges valid, and do visitors need an escort?",
                ["72|seventy-two", "escort"],
                [],
                False,
            ),
            ("When does the laboratory open on weekdays?", ["8|eight"], [], False),
            ("What is the Sunday laboratory opening time?", [], [], True),
            ("When does it open?", [], [], True),
        ],
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="docs/prehosting-results.json")
    args = parser.parse_args()
    path = ROOT / ".runtime/nist-sp1300.pdf"
    if not path.exists():
        response = httpx.get(PUBLIC_URL, timeout=60, follow_redirects=True)
        response.raise_for_status()
        path.parent.mkdir(exist_ok=True)
        path.write_bytes(response.content)
    pdf = path.read_bytes()
    if hashlib.sha256(pdf).hexdigest() != PUBLIC_SHA:
        raise SystemExit("Public fixture checksum changed; review before updating.")
    rows = []
    with httpx.Client(base_url="http://127.0.0.1:8000", timeout=190, trust_env=False) as client:
        for group, files, cases in groups(pdf):
            response = client.post("/api/sessions")
            response.raise_for_status()
            client.headers["Authorization"] = "Bearer " + response.json()["token"]
            try:
                for name, data in files:
                    response = client.post("/api/documents", files={"file": (name, data)})
                    response.raise_for_status()
                for question, patterns, locations, may_abstain in cases:
                    started = time.perf_counter()
                    response = client.post("/api/questions", json={"question": question})
                    response.raise_for_status()
                    result = response.json()
                    generated = (
                        result["mode"] == "generated"
                        and bool(patterns)
                        and all(re.search(p, result["answer"], re.I) for p in patterns)
                        and all(
                            any(s["location"] == loc for s in result["sources"])
                            for loc in locations
                        )
                    )
                    safe_refusal = may_abstain and result["mode"] == "abstained"
                    passed = bool(generated or safe_refusal)
                    rows.append(
                        {
                            "group": group,
                            "question": question,
                            "passed": passed,
                            "expected_patterns": patterns,
                            "expected_locations": locations,
                            "may_abstain": may_abstain,
                            "seconds": round(time.perf_counter() - started, 2),
                            **result,
                        }
                    )
                    print(
                        ("PASS" if passed else "REVIEW")
                        + " "
                        + question
                        + " => "
                        + result["answer"],
                        flush=True,
                    )
            finally:
                client.delete("/api/documents").raise_for_status()
    report = {
        "public_fixture": {"url": PUBLIC_URL, "sha256": PUBLIC_SHA, "pages": 9},
        "note": "Pattern/location checks require human review; safe refusal is allowed only on flagged cases. No general accuracy claim.",
        "cases": len(rows),
        "passed": sum(r["passed"] for r in rows),
        "results": rows,
    }
    (ROOT / args.output).write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    if report["cases"] != report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
